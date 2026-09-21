import uuid
import urllib.parse
import csv
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from decorators.admin import is_admin
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from django.db.models import Count, Q
from datetime import datetime
from app.models import (
    Exam,
    ExamAttempt,
    Question,
    AnswerOption,
    AttemptResponse,
)
from django.core.paginator import Paginator
import math
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.utils.timezone import localtime


# -------------------------------------------------------------------
# Vista principal: Listar exámenes disponibles para estudiante
# -------------------------------------------------------------------
@login_required(login_url="auth_login")
def available_exams(request):
    """
    Muestra los exámenes disponibles para el estudiante actual.
    """
    now = timezone.now()

    # Filtrar exámenes activos y dentro del rango de fechas
    exams = Exam.objects.filter(
        is_active=True,
        deleted_at__isnull=True,
        start_date__lte=now,
        end_date__gte=now,
        institution=request.user.institution,
    ).annotate(
        total_questions=Count(
            "question_banks__questions",
            filter=Q(
                question_banks__questions__is_active=True,
                question_banks__questions__deleted_at__isnull=True,
            ),
        )
    )

    # Para cada examen, obtener información de intentos del estudiante
    exam_data = []
    for exam in exams:
        # Ordenamos por start_time descendente para que el primero sea el más reciente
        attempts = ExamAttempt.objects.filter(user=request.user, exam=exam).order_by(
            "-start_time"
        )

        attempts_used = attempts.count()
        attempts_remaining = exam.max_attempts - attempts_used
        can_attempt = attempts_remaining > 0

        # Verificar si hay un intento en progreso usando StatusChoices
        in_progress_attempt = attempts.filter(
            status=ExamAttempt.StatusChoices.IN_PROGRESS
        ).first()

        exam_data.append(
            {
                "exam": exam,
                "attempts_used": attempts_used,
                "attempts_remaining": attempts_remaining,
                "can_attempt": can_attempt,
                "in_progress_attempt": in_progress_attempt,
                "last_attempt": attempts.first() if attempts.exists() else None,
                "total_attempts": attempts if attempts.exists() else None,
            }
        )

    paginator = Paginator(exam_data, 5)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "page_obj": page_obj,
    }

    return render(request, "exams/available_exams.html", context)


# -------------------------------------------------------------------
# Ver historial de todos los intentos del estudiante
# -------------------------------------------------------------------
@login_required(login_url="auth_login")
def my_attempts(request, student_id, exam_id=None):
    """
    Muestra el historial de todos los intentos del estudiante.
    """
    # Cambiamos student__id por user__id y started_at por start_time
    attempts = (
        ExamAttempt.objects.filter(user__id=student_id)
        .select_related("exam", "exam__institution")
        .order_by("-start_time")
    )

    if exam_id:
        attempts = attempts.filter(exam__id=exam_id)

    paginator = Paginator(attempts, 5)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {"page_obj": page_obj}

    return render(request, "exams/my_attempts.html", context)


# -------------------------------------------------------------------
# Mostrar Estudiantes que han realizado un examen específico
# -------------------------------------------------------------------
@login_required(login_url="auth_login")
@is_admin
def exam_students(request, exam_id):
    """
    Muestra los estudiantes que han realizado un examen específico.
    """
    exam = get_object_or_404(Exam, pk=exam_id, deleted_at__isnull=True)

    # Obtener intentos del examen usando "user" en lugar de "student"
    attempts = ExamAttempt.objects.filter(exam=exam).select_related("user")

    # Agrupar por estudiante (user) y contar intentos
    student_data = {}
    for attempt in attempts:
        user_id = attempt.user.id
        if user_id not in student_data:
            student_data[user_id] = {
                "student": attempt.user, # Mantenemos la key "student" para no romper tu HTML
                "attempts": [],
            }
        student_data[user_id]["attempts"].append(attempt)

    # Preparar lista para la plantilla
    exam_data = []
    for data in student_data.values():
        attempts_list = data["attempts"]
        # Ordenamos los intentos para que el último (índice 0) sea el más reciente
        attempts_list.sort(key=lambda x: x.start_time, reverse=True)
        
        attempts_used = len(attempts_list)
        attempts_remaining = exam.max_attempts - attempts_used
        can_attempt = attempts_remaining > 0

        # Verificar si hay un intento en progreso
        in_progress_attempt = next(
            (a for a in attempts_list if a.status == ExamAttempt.StatusChoices.IN_PROGRESS), None
        )

        # Calcular el porcentaje de aciertos dinámico según el algoritmo adaptativo
        total_accuracy = 0
        for a in attempts_list:
            if a.questions_answered > 0:
                accuracy = (a.correct_count / a.questions_answered) * 100
                total_accuracy += accuracy
                
        average_score = total_accuracy / attempts_used if attempts_used > 0 else 0

        exam_data.append(
            {
                "student": data["student"],
                "attempts_used": attempts_used,
                "attempts_remaining": attempts_remaining,
                "can_attempt": can_attempt,
                "in_progress_attempt": in_progress_attempt,
                "last_attempt": attempts_list[0] if attempts_list else None,
                "total_attempts": attempts_list if attempts_list else None,
                "average_score": round(average_score, 1), # Redondeado a 1 decimal
            }
        )

    paginator = Paginator(exam_data, 5)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "page_obj": page_obj,
        "exam": exam,
    }

    return render(request, "exams/exam_students.html", context)

# -------------------------------------------------------------------
# Generar y descargar reporte CSV de un intento
# -------------------------------------------------------------------
@login_required(login_url="auth_login")
@is_admin
def export_attempt_csv(request, attempt_id):
    """
    Exporta los resultados de un intento a CSV.
    """
    attempt = get_object_or_404(ExamAttempt, pk=attempt_id)

    # Crear respuesta HTTP
    response = HttpResponse(content_type="text/csv")
    filename = f"Intento # {attempt.id} de {attempt.student} ({attempt.started_at.strftime('%Y-%m-%d')}).csv"
    ascii_filename = urllib.parse.quote(filename)
    response["Content-Disposition"] = f'attachment; filename="{ascii_filename}"'
    response.write("\ufeff".encode("utf8"))

    # Generar CSV
    writer = csv.writer(response)

    # Encabezados
    writer.writerow(
        [
            "#",
            "Nivel",
            "Área",
            "Tema de la Pregunta",
            "Respuesta Seleccionada",
            "Es Correcta",
            "Tiempo (s)",
            "Tiempo Permitido (s)",
            "Violación Tiempo",
            "Índice S",
        ]
    )

    # Datos
    for answer in attempt.answers.all().order_by("question_number"):
        writer.writerow(
            [
                answer.question_number,
                answer.difficulty_level.name if answer.difficulty_level else "N/A",
                (
                    answer.question.knowledge_area.name
                    if answer.question.knowledge_area
                    else "N/A"
                ),
                (answer.question.topic if answer.question.topic else "N/A"),
                (
                    answer.selected_option.option_text[:30] + "..."
                    if answer.selected_option.option_text
                    else "Respuesta no disponible"
                ),
                "Sí" if answer.is_correct else "No",
                answer.time_taken_seconds,
                answer.allowed_time_seconds,
                "Sí" if answer.time_violation else "No",
                round(answer.s_index_after, 4),
            ]
        )

    return response


# -------------------------------------------------------------------
# Exportar reporte completo de un examen
# -------------------------------------------------------------------
@login_required(login_url="auth_login")
@is_admin
def export_exam_results(request, exam_id):
    """
    Exporta todos los resultados de un examen a CSV.
    """
    exam = get_object_or_404(Exam, pk=exam_id, deleted_at__isnull=True)

    attempts = (
        ExamAttempt.objects.filter(exam=exam)
        .exclude(status=ExamAttempt.Status.IN_PROGRESS)
        .select_related("student")
        .order_by("-started_at")
    )

    # Crear CSV
    response = HttpResponse(content_type="text/csv")
    filename = f"Resultados de {exam.title} ({timezone.now().strftime('%Y-%m-%d')}).csv"
    ascii_filename = urllib.parse.quote(filename)
    response["Content-Disposition"] = f'attachment; filename="{ascii_filename}"'
    response.write("\ufeff".encode("utf8"))

    writer = csv.writer(response)

    # Encabezados
    writer.writerow(
        [
            "#",
            "Documento",
            "Estudiante",
            "Institución",
            "Programa",
            "Examen",
            "Intento #",
            "Fecha Inicio",
            "Fecha Fin",
            "Total Preguntas",
            "Correctas",
            "Incorrectas",
            "Precisión %",
            "Índice S",
            "Resultado",
            "Consistencia",
        ]
    )

    contador = 1

    # Datos
    for attempt in attempts:
        writer.writerow(
            [
                contador,
                attempt.student.document_number,
                attempt.student or attempt.student.username,
                attempt.student.institution.name if attempt.student.institution else "N/A",
                (
                    attempt.student.academic_department.name
                    if attempt.student and attempt.student.academic_department
                    else "N/A"
                ),
                attempt.exam.title,
                attempt.attempt_number,
                attempt.started_at.strftime("%Y-%m-%d %H:%M"),
                (
                    attempt.completed_at.strftime("%Y-%m-%d %H:%M")
                    if attempt.completed_at
                    else "N/A"
                ),
                attempt.total_questions,
                attempt.correct_answers,
                attempt.incorrect_answers,
                round(attempt.get_accuracy(), 2),
                round(attempt.s_index, 4),
                attempt.get_status_display(),
                attempt.get_consistency_feedback_display() or "N/A",
            ]
        )
        contador += 1

    return response


# -------------------------------------------------------------------
# Iniciar un intento de examen
# -------------------------------------------------------------------
@login_required
def start_exam(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id)
    
    # Buscamos si hay un intento "En Progreso" para este usuario y examen
    attempt = ExamAttempt.objects.filter(
        user=request.user, 
        exam=exam, 
        status=ExamAttempt.StatusChoices.IN_PROGRESS
    ).first()
    
    # Si no existe, creamos uno nuevo
    if not attempt:
        attempt = ExamAttempt.objects.create(
            user=request.user,
            exam=exam,
            status=ExamAttempt.StatusChoices.IN_PROGRESS,
            s_history=[] # Inicializamos la lista vacía para el JSONField
        )
        
    return redirect('exam_take', attempt_id=attempt.id)


# -------------------------------------------------------------------
# Tomar un intento de examen (mostrar la siguiente pregunta)
# -------------------------------------------------------------------
@login_required
def take_exam(request, attempt_id):
    attempt = get_object_or_404(ExamAttempt, pk=attempt_id, user=request.user)
    
    # Si el intento ya no está en progreso, lo enviamos al resumen
    if attempt.status != ExamAttempt.StatusChoices.IN_PROGRESS:
        return redirect('exam_summary', attempt_id=attempt.id) # Crearemos esta ruta luego
        
    # 1. Obtenemos los IDs de las preguntas ya respondidas en este intento
    answered_question_ids = attempt.responses.values_list('question_id', flat=True)
    
    # 2. Lógica del Ascensor: Determinar el nivel de dificultad objetivo
    target_difficulty = "Básico" # Nivel por defecto al iniciar
    
    # Buscamos la última respuesta del estudiante en este intento
    last_response = attempt.responses.select_related('question__difficulty_level').order_by('-created_at').first()

    if last_response:
        # Extraemos el nombre del nivel (usamos getattr por si alguna pregunta no tiene nivel asignado)
        last_difficulty = getattr(last_response.question.difficulty_level, 'name', 'Básico')
        
        if last_response.is_correct:
            # Si acierta, sube o se mantiene en el tope
            if last_difficulty == "Básico":
                target_difficulty = "Intermedio"
            elif last_difficulty in ["Intermedio", "Avanzado"]:
                target_difficulty = "Avanzado"
        else:
            # Si falla, baja o se mantiene en el fondo
            if last_difficulty == "Avanzado":
                target_difficulty = "Intermedio"
            elif last_difficulty in ["Intermedio", "Básico"]:
                target_difficulty = "Básico"

    # 3. Plan de Contingencia (Fallback)
    # Si se agotan las preguntas del nivel objetivo, el sistema busca en el siguiente nivel más cercano
    niveles_fallback = {
        "Avanzado": ["Avanzado", "Intermedio", "Básico"],
        "Intermedio": ["Intermedio", "Básico", "Avanzado"],
        "Básico": ["Básico", "Intermedio", "Avanzado"]
    }

    next_question = None

    # 4. Buscamos la siguiente pregunta aplicando el nivel y el fallback
    for nivel in niveles_fallback.get(target_difficulty, ["Básico"]):
        next_question = Question.objects.filter(
            bank__exams=attempt.exam,
            difficulty_level__name=nivel,
            is_active=True
        ).exclude(id__in=answered_question_ids).order_by('?').first()
        
        if next_question:
            break # Si encontramos una pregunta disponible, rompemos el ciclo y la entregamos
    
    if not next_question:
        # Si ya no hay más preguntas disponibles en el banco antes de que el 
        # algoritmo decida, forzamos el fallo o cierre por falta de datos.
        attempt.status = ExamAttempt.StatusChoices.FAILED
        attempt.save()
        return redirect('exam_summary', attempt_id=attempt.id)
        
    context = {
        'attempt': attempt,
        'question': next_question,
        'exam': attempt.exam
    }
    
    return render(request, 'attempts/take_question.html', context)


# -------------------------------------------------------------------
# Procesar la respuesta enviada por el estudiante
# -------------------------------------------------------------------
@login_required
def process_answer(request, attempt_id):
    if request.method != 'POST':
        return redirect('exam_take', attempt_id=attempt_id)
        
    attempt = get_object_or_404(ExamAttempt, pk=attempt_id, user=request.user)
    
    if attempt.status != ExamAttempt.StatusChoices.IN_PROGRESS:
        return redirect('exam_summary', attempt_id=attempt.id)

    question_id = request.POST.get('question_id')
    option_id = request.POST.get('option_id') # Si viene vacío, fue timeout
    time_taken = request.POST.get('time_taken', 0)
    
    question = get_object_or_404(Question, pk=question_id)
    selected_option = None
    is_correct = False
    
    # Evaluamos si respondió algo o se le acabó el tiempo
    if option_id:
        selected_option = get_object_or_404(AnswerOption, pk=option_id)
        is_correct = selected_option.is_correct
        
    # Guardamos la respuesta en la base de datos (vital para no perder progreso)
    AttemptResponse.objects.create(
        attempt=attempt,
        question=question,
        selected_option=selected_option,
        is_correct=is_correct,
        time_taken=int(time_taken)
    )
    
    # --- INICIO DEL ALGORITMO ADAPTATIVO SECUENCIAL ---
    attempt.questions_answered += 1
    
    if is_correct:
        attempt.correct_count += 1
    else:
        attempt.incorrect_count += 1
        
    p0 = attempt.exam.p0
    p1 = attempt.exam.p1
    a = attempt.exam.a
    b = attempt.exam.b
    m = attempt.exam.max_questions
    
    c = attempt.correct_count
    w = attempt.incorrect_count
    
    # Cálculo de S (Evitamos división por cero asegurando que p0 != p1 y p0 != 100 en validaciones previas)
    s = (c * math.log(p1 / p0)) + (w * math.log((100 - p1) / (100 - p0)))
    attempt.current_s_index = s
    
    # Actualizamos el historial (JSONField)
    history = attempt.s_history
    history.append(s)
    attempt.s_history = history
    
    # Cálculo de límites
    limite_inferior = math.log(b / (1 - a))
    limite_superior = math.log((1 - b) / a)
    
    # Evaluación de estado corregida
    if s <= limite_inferior:
        # Aprobación temprana por certeza estadística
        attempt.status = ExamAttempt.StatusChoices.PASSED
        attempt.end_time = timezone.now()
    elif s >= limite_superior:
        # Reprobación temprana por certeza estadística
        attempt.status = ExamAttempt.StatusChoices.FAILED
        attempt.end_time = timezone.now()
    elif attempt.questions_answered >= m:
        # Se agotan las preguntas sin tocar límites: Evaluamos el saldo final
        # Si S <= 0, los aciertos (ponderados) superan o igualan el umbral exigido
        if s <= 0:
            attempt.status = ExamAttempt.StatusChoices.PASSED
        else:
            attempt.status = ExamAttempt.StatusChoices.FAILED
        attempt.end_time = timezone.now()
        
    attempt.save()
    
    # En lugar de ir a la siguiente pregunta de una vez, vamos a la retroalimentación
    # Pasamos el ID de la respuesta recién creada para mostrar el feedback específico
    return redirect('exam_feedback', attempt_id=attempt.id, response_id=attempt.responses.last().id)


# --------------------------------------------------------------------
# Vista de retroalimentación para una pregunta respondida
# --------------------------------------------------------------------
@login_required
def exam_feedback(request, attempt_id, response_id):
    attempt = get_object_or_404(ExamAttempt, pk=attempt_id, user=request.user)
    response = get_object_or_404(AttemptResponse, pk=response_id, attempt=attempt)
    
    # Determinamos hacia dónde debe ir el usuario al hacer clic en "Continuar"
    if attempt.status == ExamAttempt.StatusChoices.IN_PROGRESS:
        next_url = 'exam_take'
    else:
        next_url = 'exam_summary'
        
    context = {
        'attempt': attempt,
        'response': response,
        'question': response.question,
        'selected_option': response.selected_option,
        'next_url': next_url
    }
    
    return render(request, 'attempts/feedback.html', context)


# --------------------------------------------------------------------
# Vista de resumen del intento de examen
# --------------------------------------------------------------------
@login_required
def exam_summary(request, attempt_id):
    attempt = get_object_or_404(ExamAttempt, pk=attempt_id, user=request.user)
    
    historial_s = attempt.s_history
    tendencia = 0
    feedback_semaforo = ""
    color_semaforo = ""
    
    # Lógica original del semáforo
    if len(historial_s) < 3:
        feedback_semaforo = "No hay suficientes datos para evaluar consistencia."
        color_semaforo = "gray" # ⚪
    else:
        for i in range(1, len(historial_s)):
            if historial_s[i] < historial_s[i - 1]:
                tendencia += 1
            elif historial_s[i] > historial_s[i - 1]:
                tendencia -= 1
                
        mitad_longitud = len(historial_s) // 2
        
        if tendencia >= mitad_longitud:
            feedback_semaforo = "Consistente positivo"
            color_semaforo = "green" # 🟢
        elif tendencia <= -mitad_longitud:
            feedback_semaforo = "Consistente negativo"
            color_semaforo = "red" # 🔴
        else:
            feedback_semaforo = "Inconsistente"
            color_semaforo = "yellow" # 🟡

    # Evitar división por cero en el cálculo de porcentajes
    total = attempt.questions_answered
    porcentaje_correctas = (attempt.correct_count / total * 100) if total > 0 else 0
    porcentaje_incorrectas = (attempt.incorrect_count / total * 100) if total > 0 else 0

    context = {
        'attempt': attempt,
        'porcentaje_correctas': round(porcentaje_correctas, 1),
        'porcentaje_incorrectas': round(porcentaje_incorrectas, 1),
        'feedback_semaforo': feedback_semaforo,
        'color_semaforo': color_semaforo,
        'indice_s_final': attempt.current_s_index
    }
    
    return render(request, 'attempts/summary.html', context)

# --------------------------------------------------------------------
# Exportar reporte completo de un examen a Excel
# --------------------------------------------------------------------
@login_required
def export_exam_report(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id)

    # 1. Filtro por roles (Admin ve todos, Estudiante ve los suyos)
    if request.user.is_admin():
        attempts = ExamAttempt.objects.filter(exam=exam).select_related(
            'user', 'user__role', 'user__institution', 
            'user__academic_department', 'user__group', 'user__document_type'
        ).order_by('-start_time')
    else:
        attempts = ExamAttempt.objects.filter(exam=exam, user=request.user).select_related(
            'user', 'user__role', 'user__institution', 
            'user__academic_department', 'user__group', 'user__document_type'
        ).order_by('-start_time')

    # Crear el libro de Excel
    wb = openpyxl.Workbook()
    
    # Estilos para las cabeceras
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid") # Indigo 600
    header_alignment = Alignment(horizontal="center", vertical="center")

    # ==========================================
    # HOJA 1: INTENTOS DEL EXAMEN
    # ==========================================
    ws_attempts = wb.active
    ws_attempts.title = "Intentos"

    # Definir la mayor cantidad de columnas posibles para los intentos
    attempts_headers = [
        "ID", "Documento", "Nombre", "Correo", 
        "Examen", "Estado del Intento", "Preguntas Respondidas (X)", "Respuestas Correctas (C)", 
        "Respuestas Incorrectas (W)", "% Aciertos", "Índice S Actual", "Historial Score (S)",
        "Fecha/Hora Inicio", "Fecha/Hora Fin", "Duración (Minutos)"
    ]

    ws_attempts.append(attempts_headers)
    
    # Aplicar estilos a la cabecera
    for cell in ws_attempts[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Diccionario para almacenar usuarios únicos para la segunda hoja
    unique_users = {}

    for attempt in attempts:
        user = attempt.user
        unique_users[user.id] = user # Guardamos el usuario para la hoja 2

        # Cálculos adicionales
        accuracy = (attempt.correct_count / attempt.questions_answered * 100) if attempt.questions_answered > 0 else 0
        
        duration = ""
        if attempt.end_time:
            diff = attempt.end_time - attempt.start_time
            duration = round(diff.total_seconds() / 60, 2) # Duración en minutos

        start_time_str = localtime(attempt.start_time).strftime('%Y-%m-%d %H:%M:%S')
        end_time_str = localtime(attempt.end_time).strftime('%Y-%m-%d %H:%M:%S') if attempt.end_time else "En progreso"
        
        s_history_str = json.dumps(attempt.s_history) if attempt.s_history else "[]"

        row = [
            attempt.id,
            user.document_number or "N/A",
            f"{user.first_name} {user.last_name or ''}".strip(),
            user.email,
            exam.title,
            attempt.get_status_display(), # Muestra la etiqueta legible (ej. "Aprobado")
            attempt.questions_answered,
            attempt.correct_count,
            attempt.incorrect_count,
            round(accuracy, 2),
            round(attempt.current_s_index, 4),
            s_history_str,
            start_time_str,
            end_time_str,
            duration
        ]
        ws_attempts.append(row)

    # ==========================================
    # HOJA 2: DATOS PERSONALES DE ESTUDIANTES
    # ==========================================
    ws_users = wb.create_sheet(title="Datos Estudiantes")
    
    users_headers = [
        "ID", "Nombres", "Apellidos", "Correo", "Género", "Rol", "Institución", 
        "Programa Académico", "Grupo", "Tipo de Documento", "Número de Documento", 
        "Teléfono", "Semestre", "Estado Cuenta", "Fecha de Registro"
    ]
    
    ws_users.append(users_headers)
    
    # Aplicar estilos a la cabecera de la hoja 2
    for cell in ws_users[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Escribir los datos de los usuarios únicos recolectados en la iteración anterior
    for user_id, user in unique_users.items():
        row = [
            user.id,
            user.first_name,
            user.last_name or "N/A",
            user.email,
            user.gender or "N/A",
            user.role.name if user.role else "N/A",
            user.institution.name if user.institution else "N/A",
            user.academic_department.name if user.academic_department else "N/A",
            user.group.name if user.group else "N/A",
            user.document_type.name if user.document_type else "N/A",
            user.document_number or "N/A",
            user.phone or "N/A",
            user.semester or "N/A",
            "Activo" if user.is_active else "Inactivo",
            localtime(user.created_at).strftime('%Y-%m-%d %H:%M:%S')
        ]
        ws_users.append(row)

    # Ajustar el ancho de las columnas automáticamente para ambas hojas
    for sheet in [ws_attempts, ws_users]:
        for col in sheet.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            sheet.column_dimensions[column_letter].width = min(adjusted_width, 50) # Máximo 50 de ancho

    # 3. Preparar la respuesta HTTP para descargar el Excel
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    
    # Nombre del archivo dinámico
    filename = f"Reporte_{exam.title.replace(' ', '_')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response